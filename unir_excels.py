from __future__ import annotations

import argparse
from contextlib import closing
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DEFAULT_OUTPUT = "excel_unido.xlsx"
DEFAULT_SHEET_NAME = "Consolidado"
DEFAULT_EXPECTED_FILES = 29
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%m/%d/%Y",
    "%Y/%m/%d",
    "%d/%m/%Y %H:%M:%S",
    "%d-%m-%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Une archivos de Excel de una carpeta en un solo libro ordenado por fecha."
        )
    )
    parser.add_argument(
        "--carpeta",
        default=Path(__file__).resolve().parent,
        type=Path,
        help="Carpeta que contiene los archivos de Excel. Por defecto usa la carpeta de la aplicación.",
    )
    parser.add_argument(
        "--salida",
        default=DEFAULT_OUTPUT,
        help=f"Nombre del archivo de salida. Por defecto: {DEFAULT_OUTPUT}.",
    )
    parser.add_argument(
        "--hoja-salida",
        default=DEFAULT_SHEET_NAME,
        help=f"Nombre de la hoja final. Por defecto: {DEFAULT_SHEET_NAME}.",
    )
    parser.add_argument(
        "--columna-fecha",
        help="Nombre de la columna que contiene la fecha. Si no se indica, se detecta automáticamente.",
    )
    parser.add_argument(
        "--cantidad-esperada",
        default=DEFAULT_EXPECTED_FILES,
        type=int,
        help=(
            "Cantidad de archivos de Excel esperada. "
            f"Por defecto: {DEFAULT_EXPECTED_FILES}. Use 0 para omitir esta validación."
        ),
    )
    return parser.parse_args()


def is_populated_row(values: tuple[Any, ...]) -> bool:
    return any(value is not None and str(value).strip() != "" for value in values)


def normalize_header(value: Any, position: int) -> str:
    text = str(value).strip() if value is not None else ""
    return text or f"Columna_{position}"


def parse_date_value(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        for date_format in DATE_FORMATS:
            try:
                return datetime.strptime(text, date_format)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            return None
    return None


def read_excel_file(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    header_row: list[str] | None = None
    rows: list[dict[str, Any]] = []

    with closing(load_workbook(path, data_only=True, read_only=True)) as workbook:
        worksheet = workbook[workbook.sheetnames[0]]
        for row_index, raw_row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if not is_populated_row(raw_row):
                continue

            if header_row is None:
                header_row = [
                    normalize_header(cell_value, position)
                    for position, cell_value in enumerate(raw_row, start=1)
                ]
                continue

            row_data = {
                header_row[position]: raw_row[position] if position < len(raw_row) else None
                for position in range(len(header_row))
            }
            row_data["__source_file"] = path.name
            row_data["__source_row"] = row_index
            rows.append(row_data)

    if header_row is None:
        raise ValueError(f"El archivo '{path.name}' no contiene filas con datos.")

    return header_row, rows


def detect_date_column(headers: list[str], rows: list[dict[str, Any]], requested: str | None) -> str:
    if requested:
        if requested not in headers:
            raise ValueError(f"La columna de fecha '{requested}' no existe en los archivos.")
        return requested

    for header in headers:
        if "fecha" in header.lower():
            return header

    best_header = ""
    best_matches = 0
    for header in headers:
        matches = 0
        for row in rows[:50]:
            if parse_date_value(row.get(header)) is not None:
                matches += 1
        if matches > best_matches:
            best_matches = matches
            best_header = header

    if best_matches == 0:
        raise ValueError(
            "No fue posible detectar una columna de fecha. Use --columna-fecha para indicarla."
        )

    return best_header


def sort_rows_by_date(rows: list[dict[str, Any]], date_column: str) -> list[dict[str, Any]]:
    sortable_rows: list[tuple[datetime, dict[str, Any]]] = []

    for row in rows:
        parsed_date = parse_date_value(row.get(date_column))
        if parsed_date is None:
            raise ValueError(
                "No se pudo interpretar la fecha en "
                f"'{row['__source_file']}', fila {row['__source_row']}, "
                f"columna '{date_column}': {row.get(date_column)!r}"
            )
        sortable_rows.append((parsed_date, row))

    sortable_rows.sort(key=lambda item: item[0])
    return [row for _, row in sortable_rows]


def autosize_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        values = [cell.value for cell in column_cells if cell.value is not None]
        max_length = max((len(str(value)) for value in values), default=0)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 40)


def add_excel_table(worksheet: Worksheet, headers: list[str], total_rows: int) -> None:
    last_column_letter = get_column_letter(len(headers))
    reference = f"A1:{last_column_letter}{total_rows}"
    table = Table(displayName="TablaConsolidada", ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def collect_excel_files(folder: Path, output_name: str) -> list[Path]:
    files = [
        path
        for path in sorted(folder.iterdir())
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and path.name != output_name
        and not path.name.startswith("~$")
    ]
    return files


def build_consolidated_excel(
    folder: Path,
    output_name: str,
    output_sheet: str,
    date_column_name: str | None,
    expected_files: int,
) -> Path:
    if not folder.exists() or not folder.is_dir():
        raise ValueError(f"La carpeta '{folder}' no existe o no es válida.")
    if expected_files < 0:
        raise ValueError("La cantidad esperada no puede ser negativa.")

    excel_files = collect_excel_files(folder, output_name)
    if expected_files > 0 and len(excel_files) != expected_files:
        raise ValueError(
            f"Se esperaban {expected_files} archivos de Excel y se encontraron {len(excel_files)}."
        )

    all_headers: list[str] = []
    merged_rows: list[dict[str, Any]] = []

    for excel_file in excel_files:
        headers, rows = read_excel_file(excel_file)
        if not all_headers:
            all_headers = headers[:]
        else:
            for header in headers:
                if header not in all_headers:
                    all_headers.append(header)
        merged_rows.extend(rows)

    if not all_headers:
        raise ValueError("No se encontraron encabezados para consolidar.")
    if not merged_rows:
        raise ValueError("No se encontraron filas de datos para consolidar.")

    date_column = detect_date_column(all_headers, merged_rows, date_column_name)
    sorted_rows = sort_rows_by_date(merged_rows, date_column)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = output_sheet
    worksheet.append(all_headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in sorted_rows:
        worksheet.append([row.get(header) for header in all_headers])

    worksheet.freeze_panes = "A2"
    add_excel_table(worksheet, all_headers, len(sorted_rows) + 1)
    autosize_columns(worksheet)

    output_path = folder / output_name
    workbook.save(output_path)
    workbook.close()
    return output_path


def main() -> None:
    args = parse_args()
    output_path = build_consolidated_excel(
        folder=args.carpeta.resolve(),
        output_name=args.salida,
        output_sheet=args.hoja_salida,
        date_column_name=args.columna_fecha,
        expected_files=args.cantidad_esperada,
    )
    print(f"Archivo generado correctamente: {output_path}")


if __name__ == "__main__":
    main()
